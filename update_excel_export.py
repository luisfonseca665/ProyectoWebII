import os

base_dir = r"c:\Users\garci\Documentos\ITSUR\7 Semestre\Programacion Web II\ProyectoWebII\catalog"

# 1. Update views.py
views_content = """from django.shortcuts import render, get_object_or_404
from django.urls import reverse_lazy
from django.http import HttpResponse
from django.views.generic import TemplateView, ListView, DetailView, CreateView, UpdateView, DeleteView
from .models import Carrera, Materia, Alumno, Profesor, Grupo, Calificacion
from django.db.models import Avg, Count
import openpyxl
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side

class HomeView(TemplateView):
    template_name = 'home.html'

    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['num_alumnos'] = Alumno.objects.count()
        context['num_profesores'] = Profesor.objects.count()
        context['num_grupos'] = Grupo.objects.count()
        context['num_carreras'] = Carrera.objects.count()
        return context

# --- CARRERA ---
class CarreraListView(ListView): model = Carrera; template_name = 'carrera_list.html'
class CarreraDetailView(DetailView): model = Carrera; template_name = 'catalog/carrera_detail.html'
class CarreraCreateView(CreateView): model = Carrera; fields = '__all__'; template_name = 'form_generico.html'; success_url = reverse_lazy('carreras')
class CarreraUpdateView(UpdateView): model = Carrera; fields = '__all__'; template_name = 'form_generico.html'; success_url = reverse_lazy('carreras')
class CarreraDeleteView(DeleteView): model = Carrera; success_url = reverse_lazy('carreras'); template_name = 'confirm_delete.html'

# --- MATERIA ---
class MateriaListView(ListView): model = Materia; template_name = 'materia_list.html'
class MateriaDetailView(DetailView): model = Materia; template_name = 'catalog/materia_detail.html'
class MateriaCreateView(CreateView): model = Materia; fields = '__all__'; template_name = 'form_generico.html'; success_url = reverse_lazy('materias')
class MateriaUpdateView(UpdateView): model = Materia; fields = '__all__'; template_name = 'form_generico.html'; success_url = reverse_lazy('materias')
class MateriaDeleteView(DeleteView): model = Materia; success_url = reverse_lazy('materias'); template_name = 'confirm_delete.html'

# --- ALUMNO ---
class AlumnoListView(ListView):
    model = Alumno
    template_name = 'alumno_list.html'
    def get_queryset(self):
        return Alumno.objects.annotate(promedio=Avg('calificacion__calificacion_final'))

class AlumnoDetailView(DetailView): 
    model = Alumno
    template_name = 'catalog/alumno_detail.html'
    def get_context_data(self, **kwargs):
        context = super().get_context_data(**kwargs)
        context['calificaciones'] = Calificacion.objects.filter(alumno=self.object).select_related('grupo__materia', 'grupo__profesor')
        return context

class AlumnoCreateView(CreateView): model = Alumno; fields = '__all__'; template_name = 'form_generico.html'; success_url = reverse_lazy('alumnos')
class AlumnoUpdateView(UpdateView): model = Alumno; fields = '__all__'; template_name = 'form_generico.html'; success_url = reverse_lazy('alumnos')
class AlumnoDeleteView(DeleteView): model = Alumno; success_url = reverse_lazy('alumnos'); template_name = 'confirm_delete.html'

# --- PROFESOR ---
class ProfesorListView(ListView): model = Profesor; template_name = 'profesor_list.html'
class ProfesorDetailView(DetailView): model = Profesor; template_name = 'catalog/profesor_detail.html'
class ProfesorCreateView(CreateView): model = Profesor; fields = '__all__'; template_name = 'form_generico.html'; success_url = reverse_lazy('profesores')
class ProfesorUpdateView(UpdateView): model = Profesor; fields = '__all__'; template_name = 'form_generico.html'; success_url = reverse_lazy('profesores')
class ProfesorDeleteView(DeleteView): model = Profesor; success_url = reverse_lazy('profesores'); template_name = 'confirm_delete.html'

# --- GRUPO ---
class GrupoListView(ListView): model = Grupo; template_name = 'grupo_list.html'
class GrupoDetailView(DetailView): model = Grupo; template_name = 'catalog/grupo_detail.html'
class GrupoCreateView(CreateView): model = Grupo; fields = '__all__'; template_name = 'form_generico.html'; success_url = reverse_lazy('grupos')
class GrupoUpdateView(UpdateView): model = Grupo; fields = '__all__'; template_name = 'form_generico.html'; success_url = reverse_lazy('grupos')
class GrupoDeleteView(DeleteView): model = Grupo; success_url = reverse_lazy('grupos'); template_name = 'confirm_delete.html'


# --- VISTA EXPORTACIÓN EXCEL ---
def exportar_kardex_excel(request, pk):
    alumno = get_object_or_404(Alumno, pk=pk)
    calificaciones = Calificacion.objects.filter(alumno=alumno).select_related('grupo__materia', 'grupo__profesor')

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = f"Kardex_{alumno.matricula}"

    font_title = Font(name="Arial", size=14, bold=True, color="1F4E78")
    font_header = Font(name="Arial", size=11, bold=True, color="FFFFFF")
    fill_header = PatternFill(start_color="1F4E78", end_color="1F4E78", fill_type="solid")
    border_thin = Border(
        left=Side(style='thin', color='D9D9D9'),
        right=Side(style='thin', color='D9D9D9'),
        top=Side(style='thin', color='D9D9D9'),
        bottom=Side(style='thin', color='D9D9D9')
    )

    # Title Header
    ws.merge_cells("A1:G1")
    ws["A1"] = "SISTEMA INTEGRAL DE CONTROL ESCOLAR - KARDEX ACADÉMICO"
    ws["A1"].font = font_title
    ws["A1"].alignment = Alignment(horizontal="center")

    # Info Section
    ws["A3"] = "Matrícula:"
    ws["B3"] = alumno.matricula
    ws["D3"] = "Nombre:"
    ws["E3"] = f"{alumno.apellidos} {alumno.nombre}"

    ws["A4"] = "Carrera:"
    ws["B4"] = alumno.carrera.nombre
    ws["D4"] = "Semestre:"
    ws["E4"] = alumno.semestre

    ws["A5"] = "Estatus:"
    ws["B5"] = alumno.get_estatus_display()

    prom = calificaciones.aggregate(Avg('calificacion_final'))['calificacion_final__avg']
    ws["D5"] = "Promedio General:"
    ws["E5"] = round(float(prom), 2) if prom else "N/A"

    for r in range(3, 6):
        ws[f"A{r}"].font = Font(bold=True)
        ws[f"D{r}"].font = Font(bold=True)

    # Table Header
    headers = ["Código", "Materia", "Créditos", "Clave Grupo", "Profesor", "Calificación Final", "Estatus"]
    row_idx = 7
    for col_idx, header in enumerate(headers, 1):
        cell = ws.cell(row=row_idx, column=col_idx, value=header)
        cell.font = font_header
        cell.fill = fill_header
        cell.alignment = Alignment(horizontal="center")

    # Data Rows
    for calif in calificaciones:
        row_idx += 1
        materia = calif.grupo.materia
        profesor = str(calif.grupo.profesor) if calif.grupo.profesor else "Sin asignar"
        calif_val = float(calif.calificacion_final) if calif.calificacion_final is not None else 0.0
        estatus_mat = "Aprobado" if calif_val >= 70 else "Reprobado"

        ws.cell(row=row_idx, column=1, value=materia.codigo).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=2, value=materia.nombre)
        ws.cell(row=row_idx, column=3, value=materia.creditos).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=4, value=calif.grupo.clave).alignment = Alignment(horizontal="center")
        ws.cell(row=row_idx, column=5, value=profesor)
        
        c_calif = ws.cell(row=row_idx, column=6, value=calif_val)
        c_calif.alignment = Alignment(horizontal="right")
        
        c_est = ws.cell(row=row_idx, column=7, value=estatus_mat)
        c_est.alignment = Alignment(horizontal="center")
        if calif_val >= 70:
            c_est.font = Font(color="006100", bold=True)
        else:
            c_est.font = Font(color="9C0006", bold=True)

        for col in range(1, 8):
            ws.cell(row=row_idx, column=col).border = border_thin

    # Adjust Column Widths
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = openpyxl.utils.get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    response = HttpResponse(content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    response['Content-Disposition'] = f'attachment; filename="Kardex_{alumno.matricula}.xlsx"'
    wb.save(response)
    return response
"""

with open(os.path.join(base_dir, 'views.py'), 'w', encoding='utf-8') as f:
    f.write(views_content)

# 2. Update urls.py
urls_content = """from django.urls import path
from . import views

urlpatterns = [
    path('', views.HomeView.as_view(), name='home'),
    
    path('carreras/', views.CarreraListView.as_view(), name='carreras'),
    path('carreras/nueva/', views.CarreraCreateView.as_view(), name='carrera-create'),
    path('carreras/<int:pk>/', views.CarreraDetailView.as_view(), name='carrera-detail'),
    path('carreras/<int:pk>/editar/', views.CarreraUpdateView.as_view(), name='carrera-update'),
    path('carreras/<int:pk>/eliminar/', views.CarreraDeleteView.as_view(), name='carrera-delete'),

    path('materias/', views.MateriaListView.as_view(), name='materias'),
    path('materias/nueva/', views.MateriaCreateView.as_view(), name='materia-create'),
    path('materias/<int:pk>/', views.MateriaDetailView.as_view(), name='materia-detail'),
    path('materias/<int:pk>/editar/', views.MateriaUpdateView.as_view(), name='materia-update'),
    path('materias/<int:pk>/eliminar/', views.MateriaDeleteView.as_view(), name='materia-delete'),

    path('alumnos/', views.AlumnoListView.as_view(), name='alumnos'),
    path('alumnos/nuevo/', views.AlumnoCreateView.as_view(), name='alumno-create'),
    path('alumnos/<str:pk>/', views.AlumnoDetailView.as_view(), name='alumno-detail'),
    path('alumnos/<str:pk>/editar/', views.AlumnoUpdateView.as_view(), name='alumno-update'),
    path('alumnos/<str:pk>/eliminar/', views.AlumnoDeleteView.as_view(), name='alumno-delete'),
    path('alumnos/<str:pk>/exportar-kardex/', views.exportar_kardex_excel, name='alumno-exportar-kardex'),

    path('profesores/', views.ProfesorListView.as_view(), name='profesores'),
    path('profesores/nuevo/', views.ProfesorCreateView.as_view(), name='profesor-create'),
    path('profesores/<str:pk>/', views.ProfesorDetailView.as_view(), name='profesor-detail'),
    path('profesores/<str:pk>/editar/', views.ProfesorUpdateView.as_view(), name='profesor-update'),
    path('profesores/<str:pk>/eliminar/', views.ProfesorDeleteView.as_view(), name='profesor-delete'),

    path('grupos/', views.GrupoListView.as_view(), name='grupos'),
    path('grupos/nuevo/', views.GrupoCreateView.as_view(), name='grupo-create'),
    path('grupos/<int:pk>/', views.GrupoDetailView.as_view(), name='grupo-detail'),
    path('grupos/<int:pk>/editar/', views.GrupoUpdateView.as_view(), name='grupo-update'),
    path('grupos/<int:pk>/eliminar/', views.GrupoDeleteView.as_view(), name='grupo-delete'),
]
"""

with open(os.path.join(base_dir, 'urls.py'), 'w', encoding='utf-8') as f:
    f.write(urls_content)

# 3. Update alumno_detail.html to include Excel button
alumno_detail = """{% extends "pagina_maestra.html" %}
{% block title %}Detalle del Alumno{% endblock %}
{% block content %}
<div class="card mb-4">
    <div class="card-header bg-primary text-white d-flex justify-content-between align-items-center">
        <h4 class="mb-0"><i class="bi bi-person-badge"></i> {{ object.apellidos }} {{ object.nombre }}</h4>
        <a href="{% url 'alumno-exportar-kardex' object.matricula %}" class="btn btn-success btn-sm">
            <i class="bi bi-file-earmark-excel-fill"></i> Exportar Kardex a Excel
        </a>
    </div>
    <div class="card-body row">
        <div class="col-md-6">
            <p><strong>Matrícula:</strong> {{ object.matricula }}</p>
            <p><strong>Carrera:</strong> {{ object.carrera.nombre }}</p>
        </div>
        <div class="col-md-6">
            <p><strong>Estatus:</strong> {{ object.get_estatus_display }}</p>
            <p><strong>Semestre:</strong> {{ object.semestre }}</p>
        </div>
    </div>
    <div class="card-footer">
        <a href="{% url 'alumno-update' object.matricula %}" class="btn btn-warning btn-sm">Editar</a>
        <a href="{% url 'alumnos' %}" class="btn btn-secondary btn-sm">Regresar a la lista</a>
    </div>
</div>

<div class="d-flex justify-content-between align-items-center mt-4 mb-2">
    <h4>Kardex / Calificaciones</h4>
    <a href="{% url 'alumno-exportar-kardex' object.matricula %}" class="btn btn-outline-success btn-sm">
        <i class="bi bi-download"></i> Descargar Reporte (.xlsx)
    </a>
</div>

<div class="table-responsive">
    <table class="table table-bordered table-striped">
        <thead class="table-dark">
            <tr>
                <th>Código</th>
                <th>Materia</th>
                <th>Grupo</th>
                <th>Profesor</th>
                <th>Calificación Final</th>
                <th>Estatus</th>
            </tr>
        </thead>
        <tbody>
            {% for cal in calificaciones %}
            <tr>
                <td>{{ cal.grupo.materia.codigo }}</td>
                <td>{{ cal.grupo.materia.nombre }}</td>
                <td>{{ cal.grupo.clave }}</td>
                <td>{{ cal.grupo.profesor|default:"-" }}</td>
                <td>
                    {% if cal.calificacion_final %}
                        {% if cal.calificacion_final < 70 %}
                            <span class="text-danger fw-bold">{{ cal.calificacion_final }}</span>
                        {% else %}
                            <span class="text-success fw-bold">{{ cal.calificacion_final }}</span>
                        {% endif %}
                    {% else %}
                        -
                    {% endif %}
                </td>
                <td>
                    {% if cal.calificacion_final %}
                        {% if cal.calificacion_final >= 70 %}
                            <span class="badge text-bg-success">Aprobado</span>
                        {% else %}
                            <span class="badge text-bg-danger">Reprobado</span>
                        {% endif %}
                    {% else %}
                        <span class="badge text-bg-secondary">Cursando</span>
                    {% endif %}
                </td>
            </tr>
            {% empty %}
            <tr>
                <td colspan="6" class="text-center">El alumno no cuenta con calificaciones registradas.</td>
            </tr>
            {% endfor %}
        </tbody>
    </table>
</div>
{% endblock %}"""

with open(os.path.join(base_dir, 'templates', 'catalog', 'alumno_detail.html'), 'w', encoding='utf-8') as f:
    f.write(alumno_detail)

print("Exportación a Excel y vista Kardex actualizados.")
